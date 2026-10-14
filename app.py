import tkinter as tk
from tkinter import simpledialog, messagebox, filedialog
from openpyxl import load_workbook
import os
import subprocess


# ============================================================
# VARIABLES GLOBALES
# ============================================================

condiciones = {}
condicion_actual = None

pocillos = {}

archivo_excel = None
valores_excel = {}

media_control_actual = None
control_sin_ruido_actual = None


# ============================================================
# COLORES
# ============================================================

COLORES = [
    "#FF6B6B",   # rojo
    "#4D96FF",   # azul
    "#6BCB77",   # verde
    "#FFD93D",   # amarillo
    "#A66CFF",   # morado
    "#FF922B",   # naranja
    "#20C997",   # turquesa
    "#F06595",   # rosa
    "#845EF7",   # violeta
    "#12B886",   # verde
    "#E8590C",   # naranja oscuro
    "#339AF0",   # azul claro
]


# ============================================================
# CREAR CONDICIÓN
# ============================================================

def nueva_condicion():

    global condicion_actual

    numero = len(condiciones) + 1

    nombre = simpledialog.askstring(
        "Nueva condición",
        f"Nombre de la condición {numero}:"
    )

    if not nombre:
        return

    color = COLORES[(numero - 1) % len(COLORES)]

    condiciones[numero] = {
        "nombre": nombre,
        "color": color,
        "pocillos": set(),
        "control": False,
        "boton": None
    }

    crear_boton_condicion(numero)

    seleccionar_condicion(numero)


# ============================================================
# CREAR BOTÓN DE CONDICIÓN
# ============================================================

def crear_boton_condicion(numero):

    datos = condiciones[numero]

    boton = tk.Button(
        frame_lista_condiciones,
        text=f"{numero}. {datos['nombre']}",
        bg=datos["color"],
        fg="black",
        command=lambda n=numero: seleccionar_condicion(n),
        width=23,
        height=2,
        relief="raised",
        bd=3
    )

    boton.pack(
        pady=4,
        padx=5,
        fill="x"
    )

    datos["boton"] = boton


# ============================================================
# SELECCIONAR CONDICIÓN
# ============================================================

def seleccionar_condicion(numero):

    global condicion_actual

    condicion_actual = numero

    for n, datos in condiciones.items():

        boton = datos.get("boton")

        if boton is not None:

            boton.config(
                relief="raised",
                bd=3
            )

    boton_actual = condiciones[numero].get("boton")

    if boton_actual is not None:

        boton_actual.config(
            relief="solid",
            bd=6
        )

    actualizar_placa()
    actualizar_estado()


# ============================================================
# SELECCIONAR CONTROL
# ============================================================

def seleccionar_control():

    global condicion_actual

    if condicion_actual is None:

        messagebox.showwarning(
            "Aviso",
            "Primero selecciona una condición."
        )

        return

    for numero in condiciones:

        condiciones[numero]["control"] = False

    condiciones[condicion_actual]["control"] = True

    actualizar_estado()


# ============================================================
# ACTUALIZAR ESTADO
# ============================================================

def actualizar_estado():

    if condicion_actual is None:

        texto_estado.set(
            "Ninguna condición seleccionada"
        )

        return

    datos = condiciones[condicion_actual]

    texto = (
        f"Condición: {datos['nombre']}\n"
        f"Pocillos seleccionados: {len(datos['pocillos'])}"
    )

    if datos["control"]:

        texto += "\n★ CONTROL"

    texto_estado.set(texto)


# ============================================================
# SELECCIONAR POCILLO
# ============================================================

def seleccionar_pocillo(fila, columna):

    if condicion_actual is None:

        messagebox.showwarning(
            "Aviso",
            "Primero crea o selecciona una condición."
        )

        return

    clave = (fila, columna)

    if clave in valores_excel:

        valor = valores_excel[clave]

        nombre_pocillo = (
            f"{chr(65 + fila)}{columna + 1}"
        )

        etiqueta_absorbancia.set(
            f"{nombre_pocillo}  →  Absorbancia: {valor}"
        )

    if clave in condiciones[condicion_actual]["pocillos"]:

        condiciones[condicion_actual]["pocillos"].remove(
            clave
        )

    else:

        for numero, datos in condiciones.items():

            if numero != condicion_actual:

                datos["pocillos"].discard(
                    clave
                )

        condiciones[condicion_actual]["pocillos"].add(
            clave
        )

    actualizar_placa()
    actualizar_estado()


# ============================================================
# ACTUALIZAR COLORES DE LA PLACA
# ============================================================

def actualizar_placa():

    for clave, circulo in pocillos.items():

        canvas_placa.itemconfig(
            circulo,
            fill="white",
            outline="#555555",
            width=1
        )

        for numero, datos in condiciones.items():

            if clave in datos["pocillos"]:

                canvas_placa.itemconfig(
                    circulo,
                    fill=datos["color"],
                    outline="#222222",
                    width=3
                )

                break


# ============================================================
# SELECCIONAR EXCEL
# ============================================================

def seleccionar_excel():

    global archivo_excel

    archivo = filedialog.askopenfilename(
        title="Seleccionar archivo Excel",
        filetypes=[
            ("Archivos Excel", "*.xlsx")
        ]
    )

    if not archivo:
        return

    archivo_excel = archivo

    etiqueta_archivo.set(
        os.path.basename(archivo)
    )

    leer_excel()


# ============================================================
# LEER EXCEL
# ============================================================

def leer_excel():

    global valores_excel

    if archivo_excel is None:

        messagebox.showwarning(
            "Aviso",
            "Primero selecciona un archivo Excel."
        )

        return

    try:

        wb = load_workbook(
            archivo_excel,
            data_only=True
        )

        ws = wb.active

        valores_excel = {}

        # C27:N34
        for fila_excel in range(27, 35):

            for columna_excel in range(3, 15):

                fila = fila_excel - 27
                columna = columna_excel - 3

                valor = ws.cell(
                    row=fila_excel,
                    column=columna_excel
                ).value

                valores_excel[
                    (fila, columna)
                ] = valor

        etiqueta_estado_excel.set(
            "✓ Datos C27:N34 cargados"
        )

    except Exception as e:

        messagebox.showerror(
            "Error",
            f"No se pudo leer el Excel:\n\n{e}"
        )


# ============================================================
# CALCULAR CONTROL
# ============================================================

def calcular_control():

    global media_control_actual
    global control_sin_ruido_actual

    condicion_control = None

    for numero, datos in condiciones.items():

        if datos["control"]:

            condicion_control = numero
            break

    if condicion_control is None:

        messagebox.showwarning(
            "Aviso",
            "Selecciona primero una condición como CONTROL."
        )

        return

    pocillos_control = condiciones[
        condicion_control
    ]["pocillos"]

    if len(pocillos_control) == 0:

        messagebox.showwarning(
            "Aviso",
            "El CONTROL no tiene pocillos seleccionados."
        )

        return

    valores_control = []

    for pocillo in pocillos_control:

        if pocillo not in valores_excel:
            continue

        valor = valores_excel[pocillo]

        if valor is not None:

            valores_control.append(
                float(valor)
            )

    if not valores_control:

        messagebox.showwarning(
            "Aviso",
            "No se encontraron valores válidos para el CONTROL."
        )

        return

    media_control_actual = (
        sum(valores_control)
        / len(valores_control)
    )

    control_sin_ruido_actual = (
        media_control_actual - 0.04
    )

    resultado_control.set(
        f"Media CONTROL: {media_control_actual:.4f}\n"
        f"CONTROL - 0.04: {control_sin_ruido_actual:.4f}"
    )


# ============================================================
# CALCULAR PORCENTAJES
# ============================================================

def calcular_porcentajes():

    if control_sin_ruido_actual is None:

        messagebox.showwarning(
            "Aviso",
            "Primero calcula el CONTROL."
        )

        return

    if control_sin_ruido_actual == 0:

        messagebox.showerror(
            "Error",
            "El CONTROL sin ruido es 0."
        )

        return

    resultados = []

    for numero, datos in condiciones.items():

        if datos["control"]:
            continue

        for pocillo in datos["pocillos"]:

            if pocillo not in valores_excel:
                continue

            absorbancia = valores_excel[pocillo]

            if absorbancia is None:
                continue

            porcentaje = (
                (float(absorbancia) - 0.04)
                * 100
                / control_sin_ruido_actual
            )

            resultados.append(
                (
                    datos["nombre"],
                    pocillo,
                    porcentaje
                )
            )

    if not resultados:

        messagebox.showwarning(
            "Aviso",
            "No hay pocillos para calcular."
        )

        return

    print("\n========== RESULTADOS ==========")

    for nombre, pocillo, porcentaje in resultados:

        print(
            f"{nombre} | "
            f"{chr(65 + pocillo[0])}{pocillo[1] + 1} | "
            f"{porcentaje:.2f}%"
        )

    print("================================\n")

    messagebox.showinfo(
        "Porcentajes calculados",
        f"Se han calculado {len(resultados)} pocillos."
    )


# ============================================================
# GUARDAR RESULTADOS SOBRE EL ARCHIVO ORIGINAL
# ============================================================

def guardar_resultados_en_excel():

    if archivo_excel is None:

        messagebox.showwarning(
            "Aviso",
            "Primero selecciona un Excel."
        )

        return

    if media_control_actual is None:

        messagebox.showwarning(
            "Aviso",
            "Primero calcula el CONTROL."
        )

        return

    if control_sin_ruido_actual is None:

        messagebox.showwarning(
            "Aviso",
            "Primero calcula el CONTROL sin ruido."
        )

        return

    try:

        # Abrir el archivo original
        wb = load_workbook(
            archivo_excel
        )

        ws = wb.active

        # ====================================================
        # CONTROL
        # ====================================================

        ws["C36"] = media_control_actual
        ws["C37"] = control_sin_ruido_actual

        # ====================================================
        # ETIQUETAS EN FILA 26
        # ====================================================

        etiquetas_por_columna = {}

        for numero, datos in condiciones.items():

            nombre = datos["nombre"]

            for fila, columna in datos["pocillos"]:

                columna_excel = columna + 3

                if columna_excel not in etiquetas_por_columna:

                    etiquetas_por_columna[
                        columna_excel
                    ] = nombre

                else:

                    nombre_actual = (
                        etiquetas_por_columna[
                            columna_excel
                        ]
                    )

                    nombres = nombre_actual.split(" / ")

                    if nombre not in nombres:

                        etiquetas_por_columna[
                            columna_excel
                        ] = (
                            nombre_actual
                            + " / "
                            + nombre
                        )

        # Escribir nombres de las condiciones
        for columna_excel, nombre in etiquetas_por_columna.items():

            ws.cell(
                row=26,
                column=columna_excel
            ).value = nombre

        # ====================================================
        # RESULTADOS
        # ====================================================
        #
        # Primera condición no CONTROL -> E
        # Segunda -> F
        # Tercera -> G
        #
        # Resultados -> fila 38 hacia abajo
        # Promedio -> fila 45
        #
        # ====================================================

        columna_salida = 5  # E

        for numero, datos in condiciones.items():

            # No escribir resultados del CONTROL
            if datos["control"]:
                continue

            valores_porcentaje = []

            for pocillo in datos["pocillos"]:

                if pocillo not in valores_excel:
                    continue

                absorbancia = valores_excel[pocillo]

                if absorbancia is None:
                    continue

                porcentaje = (
                    (float(absorbancia) - 0.04)
                    * 100
                    / control_sin_ruido_actual
                )

                valores_porcentaje.append(
                    porcentaje
                )

            if not valores_porcentaje:

                columna_salida += 1
                continue

            # Nombre de la condición en fila 35
            ws.cell(
                row=35,
                column=columna_salida
            ).value = datos["nombre"]

            # Resultados desde fila 38
            fila_salida = 38

            for porcentaje in valores_porcentaje:

                ws.cell(
                    row=fila_salida,
                    column=columna_salida
                ).value = porcentaje

                fila_salida += 1

            # Promedio en fila 45
            promedio = (
                sum(valores_porcentaje)
                / len(valores_porcentaje)
            )

            ws.cell(
                row=45,
                column=columna_salida
            ).value = promedio

            columna_salida += 1

        # ====================================================
        # GUARDAR SOBRE EL ARCHIVO ORIGINAL
        # ====================================================

        wb.save(
            archivo_excel
        )

        # Abrir automáticamente el Excel
        subprocess.run(
            ["open", archivo_excel]
        )

        messagebox.showinfo(
            "Análisis terminado",
            "El análisis se ha guardado correctamente "
            "sobre el archivo original.\n\n"
            f"{archivo_excel}"
        )

    except Exception as e:

        messagebox.showerror(
            "Error",
            f"No se pudo guardar el archivo:\n\n{e}"
        )


# ============================================================
# VENTANA PRINCIPAL
# ============================================================

ventana = tk.Tk()

ventana.title(
    "MTT Analyzer"
)

ventana.geometry(
    "1400x800"
)

ventana.minsize(
    1200,
    700
)


# ============================================================
# PANEL IZQUIERDO
# ============================================================

panel_izquierdo = tk.Frame(
    ventana,
    width=280
)

panel_izquierdo.pack(
    side="left",
    fill="y",
    padx=10,
    pady=10
)

panel_izquierdo.pack_propagate(
    False
)


tk.Label(
    panel_izquierdo,
    text="CONDICIONES",
    font=("Arial", 16, "bold")
).pack(
    pady=(5, 10)
)


tk.Button(
    panel_izquierdo,
    text="+ Nueva condición",
    command=nueva_condicion,
    width=25,
    height=2
).pack(
    pady=4
)


tk.Button(
    panel_izquierdo,
    text="★ Seleccionar CONTROL",
    command=seleccionar_control,
    width=25,
    height=2
).pack(
    pady=4
)


# ============================================================
# LISTA DE CONDICIONES
# ============================================================

frame_scroll = tk.Frame(
    panel_izquierdo
)

frame_scroll.pack(
    fill="both",
    expand=True,
    pady=8
)


canvas_condiciones = tk.Canvas(
    frame_scroll,
    highlightthickness=0
)

scrollbar = tk.Scrollbar(
    frame_scroll,
    orient="vertical",
    command=canvas_condiciones.yview
)

canvas_condiciones.configure(
    yscrollcommand=scrollbar.set
)

scrollbar.pack(
    side="right",
    fill="y"
)

canvas_condiciones.pack(
    side="left",
    fill="both",
    expand=True
)


frame_lista_condiciones = tk.Frame(
    canvas_condiciones
)

canvas_condiciones.create_window(
    (0, 0),
    window=frame_lista_condiciones,
    anchor="nw",
    width=245
)


def actualizar_scroll(event=None):

    canvas_condiciones.configure(
        scrollregion=canvas_condiciones.bbox("all")
    )


frame_lista_condiciones.bind(
    "<Configure>",
    actualizar_scroll
)


# ============================================================
# ESTADO
# ============================================================

texto_estado = tk.StringVar()

texto_estado.set(
    "Ninguna condición seleccionada"
)

tk.Label(
    panel_izquierdo,
    textvariable=texto_estado,
    justify="left",
    anchor="w",
    relief="groove",
    padx=8,
    pady=8
).pack(
    fill="x",
    pady=5
)


# ============================================================
# PANEL CENTRAL
# ============================================================

panel_central = tk.Frame(
    ventana
)

panel_central.pack(
    side="left",
    fill="both",
    expand=True,
    padx=15,
    pady=10
)


tk.Label(
    panel_central,
    text="PLACA DE 96 POCILLOS",
    font=("Arial", 18, "bold")
).pack(
    pady=10
)


frame_placa = tk.Frame(
    panel_central
)

frame_placa.pack(
    pady=10
)


# ============================================================
# CANVAS DE LA PLACA
# ============================================================

canvas_placa = tk.Canvas(
    frame_placa,
    width=700,
    height=500,
    bg="white",
    highlightthickness=1,
    highlightbackground="#BBBBBB"
)

canvas_placa.pack()


# ============================================================
# DIMENSIONES DE LOS POCILLOS
# ============================================================

radio = 16

separacion_x = 52
separacion_y = 52

inicio_x = 55
inicio_y = 50


# ============================================================
# NÚMEROS 1-12
# ============================================================

for columna in range(12):

    x = inicio_x + columna * separacion_x

    canvas_placa.create_text(
        x,
        15,
        text=str(columna + 1),
        font=("Arial", 10, "bold")
    )


# ============================================================
# LETRAS A-H
# ============================================================

for fila in range(8):

    y = inicio_y + fila * separacion_y

    canvas_placa.create_text(
        15,
        y,
        text=chr(65 + fila),
        font=("Arial", 10, "bold")
    )


# ============================================================
# CREAR LOS 96 POCILLOS
# ============================================================

for fila in range(8):

    for columna in range(12):

        x = inicio_x + columna * separacion_x
        y = inicio_y + fila * separacion_y

        circulo = canvas_placa.create_oval(
            x - radio,
            y - radio,
            x + radio,
            y + radio,
            fill="white",
            outline="#555555",
            width=1
        )

        pocillos[
            (fila, columna)
        ] = circulo

        canvas_placa.tag_bind(
            circulo,
            "<Button-1>",
            lambda event, f=fila, c=columna:
                seleccionar_pocillo(f, c)
        )


# ============================================================
# ABSORBANCIA
# ============================================================

etiqueta_absorbancia = tk.StringVar()

etiqueta_absorbancia.set(
    "Selecciona un pocillo"
)

tk.Label(
    panel_central,
    textvariable=etiqueta_absorbancia,
    font=("Arial", 12),
    relief="groove",
    padx=10,
    pady=8
).pack(
    fill="x",
    pady=15
)


# ============================================================
# PANEL DERECHO
# ============================================================

panel_derecho = tk.Frame(
    ventana,
    width=300
)

panel_derecho.pack(
    side="right",
    fill="y",
    padx=10,
    pady=10
)

panel_derecho.pack_propagate(
    False
)


tk.Label(
    panel_derecho,
    text="ANÁLISIS",
    font=("Arial", 16, "bold")
).pack(
    pady=(5, 15)
)


# ============================================================
# EXCEL
# ============================================================

tk.Label(
    panel_derecho,
    text="ARCHIVO EXCEL",
    font=("Arial", 11, "bold")
).pack(
    pady=5
)


tk.Button(
    panel_derecho,
    text="📂 Seleccionar Excel",
    command=seleccionar_excel,
    width=27,
    height=2
).pack(
    pady=5
)


etiqueta_archivo = tk.StringVar()

etiqueta_archivo.set(
    "Ningún Excel seleccionado"
)

tk.Label(
    panel_derecho,
    textvariable=etiqueta_archivo,
    wraplength=250,
    justify="center"
).pack(
    pady=5
)


etiqueta_estado_excel = tk.StringVar()

etiqueta_estado_excel.set(
    ""
)

tk.Label(
    panel_derecho,
    textvariable=etiqueta_estado_excel
).pack(
    pady=5
)


# ============================================================
# CONTROL
# ============================================================

tk.Label(
    panel_derecho,
    text="CONTROL",
    font=("Arial", 11, "bold")
).pack(
    pady=(20, 5)
)


tk.Button(
    panel_derecho,
    text="🧮 Calcular CONTROL",
    command=calcular_control,
    width=27,
    height=2
).pack(
    pady=5
)


resultado_control = tk.StringVar()

resultado_control.set(
    "CONTROL todavía no calculado"
)

tk.Label(
    panel_derecho,
    textvariable=resultado_control,
    justify="left",
    relief="groove",
    padx=8,
    pady=8
).pack(
    fill="x",
    pady=8
)


# ============================================================
# PORCENTAJES
# ============================================================

tk.Label(
    panel_derecho,
    text="RESULTADOS",
    font=("Arial", 11, "bold")
).pack(
    pady=(15, 5)
)


tk.Button(
    panel_derecho,
    text="📊 Calcular porcentajes",
    command=calcular_porcentajes,
    width=27,
    height=2
).pack(
    pady=5
)


# ============================================================
# GUARDAR
# ============================================================

tk.Button(
    panel_derecho,
    text="💾 GUARDAR RESULTADOS",
    command=guardar_resultados_en_excel,
    width=27,
    height=2
).pack(
    pady=(20, 5)
)


tk.Label(
    panel_derecho,
    text="Los resultados se guardarán\n"
         "sobre el archivo original.",
    justify="center"
).pack(
    pady=5
)


# ============================================================
# INICIAR
# ============================================================

ventana.mainloop()